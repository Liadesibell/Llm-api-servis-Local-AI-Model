from __future__ import annotations

from pathlib import Path

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_DOCUMENT_CHARS = 6_000
SUPPORTED_EXTENSIONS = {
    ".c",
    ".cpp",
    ".css",
    ".csv",
    ".docx",
    ".go",
    ".h",
    ".html",
    ".java",
    ".js",
    ".json",
    ".jsx",
    ".md",
    ".pdf",
    ".php",
    ".py",
    ".rb",
    ".rs",
    ".sql",
    ".ts",
    ".tsx",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}


def read_document(
    file_path: Path,
    max_chars: int = MAX_DOCUMENT_CHARS,
) -> str:
    extension = file_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Desteklenmeyen dosya türü: {extension or 'uzantısız'}")
    if file_path.stat().st_size > MAX_FILE_SIZE:
        raise ValueError("Dosya 10 MB sınırını aşıyor.")

    if extension == ".pdf":
        from pypdf import PdfReader

        content = "\n".join(page.extract_text() or "" for page in PdfReader(file_path).pages)
    elif extension == ".docx":
        from docx import Document

        document = Document(str(file_path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(
                " | ".join(cell.text for cell in row.cells) for row in table.rows
            )
        content = "\n".join(paragraphs)
    else:
        content = file_path.read_text(encoding="utf-8-sig")

    if not content.strip():
        raise ValueError("Dosyada okunabilir metin bulunamadı.")
    if len(content) > max_chars:
        formatted_limit = f"{max_chars:,}".replace(",", ".")
        raise ValueError(
            f"Dosyadan çıkarılan metin {formatted_limit} karakter sınırını aşıyor. "
            "Daha küçük veya ilgili bölümü içeren bir dosya seç."
        )
    return content
